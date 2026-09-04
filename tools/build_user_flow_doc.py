from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "跟一口_User Flow_V0.1.docx"

# Resolved preset: compact_reference_guide.
# First-page header pattern: memo_masthead.
# Named overrides:
# - CJKFont: Microsoft YaHei for reliable Chinese rendering.
# - MastheadTitle: 27 pt title and a restrained teal accent for project identity.
# - FlowLabel: compact shaded labels for Trigger / Preconditions / End state.
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1)
HEADER_DISTANCE = Inches(0.492)
FOOTER_DISTANCE = Inches(0.492)
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
INK = RGBColor(0x1B, 0x2A, 0x30)
BLACK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x68, 0x76, 0x7C)
TEAL = RGBColor(0x2E, 0x7D, 0x79)
DEEP_TEAL = RGBColor(0x1F, 0x5F, 0x5B)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
LIGHT_BLUE_FILL = "E8EEF5"
LIGHT_TEAL_FILL = "EAF4F2"
LIGHT_GRAY_FILL = "F4F6F7"
RULE = "B8C8CC"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


FLOW_SUMMARY = [
    ("UF-01", "核心喝水闭环", "日常主路径", "从“我喝了”到“跟一口”的完整闭环，回应不再触发新提醒。"),
    ("UF-02", "首次创建、加入与扩展关系", "关系建立", "支持 1 对 1 与多人小组；新增成员须经全部现有成员同意。"),
    ("UF-03", "多信号汇总与选择性回应", "日常分支", "最新信号突出、多人汇入当前轮次；默认回应全部，也可只回应部分。"),
    ("UF-04", "信号可见性与重叠关系", "关系控制", "接收与发送控制相互独立；同一信号跨重叠关系只呈现一次。"),
    ("UF-05", "退出关系与成员权限", "关系退出", "成员可以自由退出，但任何人都不能把别人踢出去。"),
    ("UF-06", "离线与通知不可用", "异常路径", "不承诺送达、不显示已读；用户返回后只看到汇总后的当前轮次。"),
]


def set_run_font(run, size=None, color=INK, bold=None, italic=None, font_name=FONT_LATIN):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="C9D4D8", size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("Table widths must sum to 9360 DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_paragraph_border(paragraph, color=RULE, size="8", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fallback, fld_end])
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part.element
    existing_abstract = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing_abstract, default=-1) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    if kind == "bullet":
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Arial")
        r_fonts.set(qn("w:hAnsi"), "Arial")
        r_pr.append(r_fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(existing_num, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return abstract_id, num_id


def add_num_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(existing_num, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.right_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("跟一口  ·  USER FLOW V0.1")
    set_run_font(hr, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)

    doc.core_properties.title = "跟一口 User Flow V0.1"
    doc.core_properties.subject = "已确认用户流程汇总"
    doc.core_properties.author = "跟一口产品项目"
    doc.core_properties.keywords = "User Flow, 产品设计, 跟一口"


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False,
             before=0, after=6, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.LEFT,
             keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_rich_para(doc, parts, before=0, after=6, line_spacing=1.25, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    for text, bold, color in parts:
        r = p.add_run(text)
        set_run_font(r, size=11, color=color, bold=bold)
    return p


def add_bullet(doc, text, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    apply_numbering(p, num_id)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_numbered_steps(doc, steps, abstract_id):
    num_id = add_num_instance(doc, abstract_id)
    for step in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        apply_numbering(p, num_id)
        r = p.add_run(step)
        set_run_font(r, size=11, color=INK)


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.20
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, size=9.5, color=DEEP_TEAL, bold=True)
    r2 = p.add_run(value)
    set_run_font(r2, size=10.5, color=INK)
    shade_paragraph(p, LIGHT_GRAY_FILL)
    return p


def add_callout(doc, label, text, fill=LIGHT_TEAL_FILL):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.22
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, size=10.5, color=DEEP_TEAL, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    shade_paragraph(p, fill)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    for r in p.runs:
        set_run_font(
            r,
            size={1: 16, 2: 13, 3: 12}[level],
            color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level],
            bold=True,
        )
    return p


def add_masthead(doc):
    add_para(doc, "PRODUCT FLOW SPEC", size=9.5, color=TEAL, bold=True, after=4)
    add_para(doc, "跟一口 · User Flow", size=27, color=INK, bold=True, after=3, line_spacing=1.0)
    add_para(doc, "已确认用户流程汇总", size=14, color=MUTED, after=15, line_spacing=1.1)

    metadata = [
        ("文档版本", "V0.1"),
        ("产品阶段", "User Flow 已完成，待进入页面／状态清单"),
        ("确认日期", "2026 年 8 月 17 日"),
        ("文档性质", "阶段性交付 · 后续可持续更新"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, size=10.5, color=INK, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(5)
    rule.paragraph_format.space_after = Pt(11)
    set_paragraph_border(rule, color="6FA7A3", size="12", space="7")

    add_callout(
        doc,
        "体验北极星",
        "产品不是监督别人喝水，而是让平静的人能够被身边人的日常行为轻轻带动一下。",
        LIGHT_TEAL_FILL,
    )
    add_para(
        doc,
        "本文件只整理已经确认的用户行为与系统反应，不提前确定页面布局、按钮形式、颜色、动效、通知载体或技术方案。流程导图可独立保存，并以本文的文字规则作为解释依据。",
        size=10.5,
        color=MUTED,
        after=4,
    )


def add_flow_overview_table(doc):
    table = doc.add_table(rows=1, cols=4)
    widths = [1040, 2500, 1300, 4520]
    headers = ["编号", "流程", "类型", "本流程解决的问题"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        shade_cell(cell, LIGHT_BLUE_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=DARK_BLUE, bold=True)
    set_repeat_table_header(table.rows[0])

    for flow_id, name, kind, purpose in FLOW_SUMMARY:
        cells = table.add_row().cells
        values = [flow_id, name, kind, purpose]
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(value)
            set_run_font(r, size=9.3, color=INK, bold=(idx == 0))
    set_table_geometry(table, widths)
    set_table_borders(table)
    return table


def build_document():
    doc = Document()
    configure_document(doc)
    bullet_abstract, bullet_num_id = add_numbering_definition(doc, "bullet")
    decimal_abstract, _ = add_numbering_definition(doc, "decimal")

    add_masthead(doc)

    add_heading(doc, "1. 文档定位与阅读方式", 1)
    add_para(
        doc,
        "这份文档承接 PRD 与《产品决策记录 Q&A》。PRD 说明产品要解决什么问题，决策记录解释关键选择为什么这样定，本文则把这些选择组织成用户实际会经历的完整流程。",
    )
    add_rich_para(
        doc,
        [
            ("已确认：", True, DEEP_TEAL),
            ("已经由产品创作者明确同意，可以进入页面／状态清单和原型设计。", False, INK),
        ],
    )
    add_rich_para(
        doc,
        [
            ("暂缓决定：", True, MUTED),
            ("不是缺失，而是应在视觉、交互或技术阶段再决定，避免越级锁死方案。", False, INK),
        ],
    )

    add_heading(doc, "2. 核心对象与全局规则", 1)
    add_heading(doc, "2.1 核心对象", 2)
    definitions = [
        ("喝水信号（Drink Signal）", "用户主动选择“我喝了”后创建的新信号。它是独立的主动表达，并向符合当前关系和可见性设置的关联成员分发。"),
        ("跟随回应（Follow Response）", "用户选择“跟一口”后，对当前轮次中的发起者给出的轻量回应。它不会再创建喝水信号，也不会形成通知接力。"),
        ("当前轮次（Current Round）", "用户尚未结束的一组来信。一个或多个人的喝水信号可以汇入同一轮次，最新信号更突出，其他来源保留为轻量背景。"),
        ("关系（Relationship）", "由 1 对 1 或少量成员共同组成的连接网络。关系名称只用于识别，不代表固定的情侣、宿舍或群聊身份。"),
        ("关系提案（Membership Proposal）", "现有成员提出新增候选人后，由全部现有成员共同确认的过程；候选人最终也必须接受邀请。"),
    ]
    for label, value in definitions:
        add_rich_para(doc, [(f"{label}：", True, DEEP_TEAL), (value, False, INK)], after=5)

    add_heading(doc, "2.2 全局规则", 2)
    global_rules = [
        "“我喝了”和“跟一口”是两种独立的信号源，系统不把其中一种自动解释成另一种。",
        "同一口水属于“我喝了”还是“跟一口”，由用户根据当下的个人判断选择；产品不追问动机，也不校验先点击还是先喝水。",
        "所有关系成员地位平等，没有群主、管理员、固定提醒者或被提醒者。",
        "没有未完成、落后、错过、连续签到、催促或惩罚；只呈现已经发生的正向行为。",
        "发起者只会看到别人主动给出的“跟一口”反馈，不会看到送达、已读、在线或未回应名单。",
        "具体提示载体、颜色、动效、声音和页面布局全部留到后续设计阶段。",
    ]
    for item in global_rules:
        add_bullet(doc, item, bullet_num_id)

    add_heading(doc, "3. User Flow 总览", 1)
    add_flow_overview_table(doc)
    add_para(
        doc,
        "说明：以下流程编号沿用本轮讨论形成的六个主题。每个流程均包含触发条件、主路径、分支与结束状态，方便后续直接映射到页面和状态。",
        size=9.5,
        color=MUTED,
        before=6,
        after=3,
    )

    doc.add_page_break()
    add_heading(doc, "4. 详细流程", 1)

    add_heading(doc, "UF-01  核心喝水闭环", 2)
    add_label_value(doc, "触发条件", "成员在现实中喝水，并决定在产品中表达这次行为。")
    add_label_value(doc, "前置条件", "用户至少存在一名可接收信号的关联成员；若尚未建立关系，则转入 UF-02。")
    add_heading(doc, "主路径", 3)
    add_numbered_steps(doc, [
        "用户根据自己的当下判断，选择“我喝了”。",
        "系统判断是否属于短时间内的重复误触：若是，则合并；若不是，则创建新的喝水信号。",
        "系统依据用户的关系网络和可见性设置，计算符合条件的接收成员；重叠关系中的同一接收者只计算一次。",
        "喝水信号被发出，发起者获得轻量的“已发出”确认。",
        "接收成员感知到一次轻提示；提示形式暂不在本阶段规定。",
        "接收成员可以选择“跟一口”，也可以忽略，或根据自己的判断独立选择“我喝了”。",
        "如果接收成员选择“跟一口”，系统创建跟随回应，并给该成员轻量成功反馈。",
        "原喝水信号的发起者收到有人跟上的正向反馈；这条反馈不会再次提醒其他成员。",
    ], decimal_abstract)
    add_heading(doc, "分支与边界", 3)
    for item in [
        "用户可以先喝水再点击，也可以先点击后喝水；系统不验证现实顺序。",
        "成员不参与“跟一口”不会产生任何负面标记，也不会收到追问或二次提醒。",
        "同一成员短时间重复点击默认合并；具体合并窗口留到技术结构阶段。",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_label_value(doc, "结束状态", "主动信号已创建；跟随回应如有发生，仅完成正向反馈，不形成新的提醒链。")

    add_heading(doc, "UF-02  首次创建、加入与扩展关系", 2)
    add_label_value(doc, "触发条件", "新用户首次使用，或现有成员希望建立一份新的关系。")
    add_heading(doc, "A. 创建或加入第一份关系", 3)
    add_numbered_steps(doc, [
        "用户完成最小身份建立；具体资料字段留到页面／状态清单阶段。",
        "用户根据自己的现实场景选择创建关系，或通过已有邀请加入关系。",
        "创建者设置一个便于识别的关系名称，并生成邀请方式；关系名称只是标签，不定义成员的真实关系类型。",
        "被邀请者查看关系与成员信息后，自主接受或拒绝。",
        "接受后双方成为平等成员；创建者不保留群主或管理权限。",
    ], decimal_abstract)
    add_heading(doc, "B. 向既有多人关系新增成员", 3)
    add_numbered_steps(doc, [
        "任意现有成员都可以提出一名候选人。",
        "系统向全部现有成员发起加入确认；所有受影响的现有成员都必须同意。",
        "只要有人拒绝或尚未确认，候选人就不会正式加入。",
        "全部现有成员同意后，候选人才收到正式邀请，并由候选人决定是否接受。",
        "候选人接受后成为平等成员，不产生邀请者特权。",
    ], decimal_abstract)
    add_heading(doc, "分支与边界", 3)
    for item in [
        "A 与 B 已有关系时，B 不能单独决定把 C 加入；因为 C 的加入也会影响 A。",
        "如果 B 想与 C 建立独立连接，可以另行创建 B-C 关系，不需要改变 A-B 关系。",
        "不同关系即使共享成员也不会自动合并，B 同时存在于 A-B 与 B-C 中，不代表 A 与 C 建立连接。",
        "全员确认会增加少量步骤，但当前关系体量较小，平等和知情优先于极致加入效率。",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_label_value(doc, "结束状态", "关系建立或成员加入完成后，所有成员拥有相同的日常发起、回应和退出权利。")

    add_heading(doc, "UF-03  多信号汇总与选择性回应", 2)
    add_label_value(doc, "触发条件", "用户尚未处理当前来信时，又有一个或多个人发出“我喝了”。")
    add_heading(doc, "主路径", 3)
    add_numbered_steps(doc, [
        "第一条来信形成当前轮次。",
        "新的喝水信号到达时，系统更新当前轮次，而不是创建多条待办。",
        "最新信号在当前轮次中更突出；此前发起者仍作为轻量来源信息保留。",
        "用户点击默认的“跟一口”时，一次真实喝水和一次操作回应当前轮次中的全部发起者。",
        "每位被回应的发起者分别收到轻量跟随反馈；其他无关成员不需要看到。",
        "用户完成操作后，本轮结束，不留下未读、未回复或稍后处理任务。",
    ], decimal_abstract)
    add_heading(doc, "选择性回应", 3)
    add_numbered_steps(doc, [
        "用户通过次级入口进入选择状态；具体是长按、显式入口还是点击成员头像，留到交互设计阶段。",
        "系统默认保持低成本操作，用户只在有需要时选择希望回应的成员。",
        "被选中的成员收到跟随反馈；未被选择的成员不收到回应，也不会知道自己被排除。",
        "选择完成后整个当前轮次结束，未选成员不会作为待回复对象继续留在首页。",
    ], decimal_abstract)
    add_heading(doc, "独立选择“我喝了”", 3)
    for item in [
        "用户可以忽略当前轮次，并根据自己的判断选择“我喝了”。",
        "新的主动信号正常发出，但不算对当前任何人的“跟一口”。",
        "此前的来信轮次随之结束；用户不会在之后补回这一轮的跟随回应。",
        "产品不判断这口水是否实际上受别人影响，只记录用户主动表达的社交意图。",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_label_value(doc, "结束状态", "当前轮次通过默认跟随、选择性跟随或独立“我喝了”结束；没有回复债务。")

    doc.add_page_break()
    add_heading(doc, "UF-04  信号可见性与重叠关系", 2)
    add_label_value(doc, "触发条件", "用户希望长期调整某个人或某一份关系的信号往来，或同一成员通过多份关系与自己相连。")
    add_heading(doc, "默认分发", 3)
    for item in [
        "“我喝了”是用户级的一次主动信号，默认分发给所有仍然具有有效连接的关联成员。",
        "常规路径不提供每次发送前的收件人选择，以保留一键表达的低成本体验。",
        "如果同一人通过多份关系与发起者相连，同一次喝水信号只呈现和提醒一次。",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "双向但相互独立的控制", 3)
    add_numbered_steps(doc, [
        "“不看 TA 的信号”控制进入自己的信息：对该人的设置在所有共同关系中生效。",
        "“不让 TA 看我的信号”控制自己的发出范围：同样可针对这个人全局生效。",
        "两个设置可以单独开启，也可以同时开启；不要求对称。",
        "针对整份关系的暂停只作用于该关系。如果双方还共享另一份未暂停关系，仍可能通过那份关系保持信号连接。",
        "所有控制静默生效，不通知对方，也不公开展示给关系中的其他成员。",
    ], decimal_abstract)
    add_heading(doc, "关系来源", 3)
    for item in [
        "系统需要保留成员与关系来源信息，便于后续界面区分不同关系中的跟随反馈。",
        "未来可以使用颜色、成员头像、名称或关系标签辅助识别，但颜色不能成为唯一的区分方式。",
        "具体颜色与视觉规则留到品牌与视觉概念阶段，由产品创作者本人决定。",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_label_value(doc, "结束状态", "用户获得可逆、非对称、不会公开伤害关系的信号控制，同时保留关系成员身份。")

    add_heading(doc, "UF-05  退出关系与成员权限", 2)
    add_label_value(doc, "触发条件", "成员希望结束自己在某份关系中的参与。")
    add_heading(doc, "主路径", 3)
    add_numbered_steps(doc, [
        "成员主动发起退出，并看到一次明确确认，避免误触。",
        "确认后无需其他成员批准，退出立即生效。",
        "其他成员收到一条中性的关系变更说明，而不是喝水提醒。",
        "多人关系剩余两人及以上时继续存在；只剩一人时自动结束。",
        "一对一关系中任意一方退出后，这段关系结束。",
    ], decimal_abstract)
    add_heading(doc, "成员移除边界", 3)
    for item in [
        "产品不提供“踢出成员”能力。任何个人或成员组合都不能替另一个人决定其成员身份。",
        "如果 A 与 B 不想继续与 C 共处，可以分别退出原关系，再建立新的 A-B 关系。",
        "这项限制避免形成隐藏的群主权力，也减少产品成为排斥或微缩霸凌工具的可能。",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_callout(doc, "权限原则", "每个人只能决定自己是否继续参与关系，不能替别人决定其成员身份。")
    add_label_value(doc, "结束状态", "退出者与原关系断开；剩余关系按人数继续或结束，不产生驱逐权限。")

    add_heading(doc, "UF-06  离线与通知不可用", 2)
    add_label_value(doc, "触发条件", "信号发出时，接收成员没联网、关闭系统通知，或手机处于免打扰状态。")
    add_heading(doc, "主路径", 3)
    add_numbered_steps(doc, [
        "发起者完成“我喝了”，系统给出“信号已发出”的结果。",
        "系统不把“已发出”描述成“所有成员已收到提醒”，也不向发起者展示逐人送达状态。",
        "接收成员没有获得即时系统通知时，不产生失败、落后或未读处罚。",
        "接收成员返回应用后，只看到汇总后的当前轮次，而不是一串需要逐条处理的过期提醒。",
        "接收成员可以正常跟随、选择部分回应、独立选择“我喝了”或继续忽略。",
        "发起者只会在对方主动“跟一口”时收到反馈，不会看到已读、在线或看过未回应。",
    ], decimal_abstract)
    add_heading(doc, "边界", 3)
    for item in [
        "通知权限关闭时，只在适当位置平静说明，不反复弹窗催促开启。",
        "不补发连续提示，不建立未读数字，也不把离线期间的信号变成任务列表。",
        "发起者自身离线时的本地保存、重试和失败恢复方式留到技术结构阶段。",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_label_value(doc, "结束状态", "提醒可能延迟，但社交关系不产生监督感；显式跟随仍是唯一可见的回应。")

    add_heading(doc, "5. 跨流程的异常与防误触规则", 1)
    cross_rules = [
        "短时间内同一成员的重复“我喝了”默认合并，用于处理误触；正常间隔后的再次喝水仍允许形成新信号。",
        "不为喝水信号设置强调紧迫感的倒计时；用户未结束当前轮次时，新信号继续更新该轮次。",
        "用户通过“跟一口”或“我喝了”完成一次明确选择后，当前轮次结束。",
        "被忽略、未选中、未读或没有通知权限都不会被解释成负面社交行为。",
        "成员和关系控制均可恢复；具体恢复入口与反馈方式留到页面／状态清单阶段。",
    ]
    for item in cross_rules:
        add_bullet(doc, item, bullet_num_id)

    add_heading(doc, "6. 已确认的产品原则", 1)
    principles = [
        ("轻影响，不监督", "产品只提供来自熟人的自然行为提示，不追踪谁没喝、谁没回。"),
        ("用户解释自己的行为", "同一现实动作可能有不同社交意图，系统记录用户选择，不推断因果。"),
        ("默认低成本，例外有选择权", "常规一次操作完成，复杂选择通过次级入口按需出现。"),
        ("关系平等，但个人边界可不对称", "没有群主和踢人权；接收与发送控制可以分别设置。"),
        ("只传播正向反馈", "“跟一口”能让发起者感到自己的行为带动了别人，但回应不继续扩散。"),
        ("没有回复债务", "多条来信汇成当前轮次，处理后整体结束，不制造未读任务。"),
    ]
    for label, value in principles:
        add_rich_para(doc, [(f"{label}：", True, DEEP_TEAL), (value, False, INK)], after=5)

    doc.add_page_break()
    add_heading(doc, "7. 暂缓到后续阶段决定", 1)
    deferred = [
        "页面／状态清单：具体页面数量、首页结构、关系管理入口、邀请审批状态、空状态和通知关闭状态。",
        "低保真 Wireframe：选择性回应的次级入口采用长按、显式文字入口还是成员点击；不在本文锁定。",
        "品牌与视觉概念：产品最终名称、关系与成员的颜色规则、动画、声音、文案语气和反馈强度。",
        "技术结构：重复点击合并窗口、发起者离线时的保存与重试、推送渠道、数据保留与同步方式。",
        "MVP 范围：移动 App、小程序或其他载体；关系人数上限和历史记录深度。",
        "真实测试：3～5 名朋友使用约一周后，再根据真实行为判断哪些规则需要进入 V0.2。",
    ]
    for item in deferred:
        add_bullet(doc, item, bullet_num_id)

    add_heading(doc, "8. 向页面／状态清单的交接", 1)
    add_para(
        doc,
        "下一阶段不再讨论“产品逻辑是否成立”，而是把每条流程拆成用户能够看见和操作的页面状态。至少需要覆盖以下场景，但此处不规定视觉形式：",
    )
    handoff = [
        "尚未建立关系；创建关系；收到邀请；等待现有成员确认；候选人接受或拒绝。",
        "没有当前来信；单人来信；多人汇总；默认跟随；选择性回应；独立“我喝了”。",
        "信号发出成功；跟随成功；发起者收到跟随反馈；通知不可用；离线或发送失败。",
        "关系列表；关系详情；针对人或关系的收发控制；退出确认；关系结束。",
    ]
    for item in handoff:
        add_bullet(doc, item, bullet_num_id)
    add_callout(doc, "下一阶段的首要问题", "用户每天打开产品时，首页第一眼应该呈现什么，以及首页如何在“平静”“来信”“回应成功”之间切换。")

    add_heading(doc, "9. 产品经理求职／作品集表达提示", 1)
    add_para(
        doc,
        "这组 User Flow 的价值不只是流程完整，还体现了产品判断。作品集中可以按照“问题—决策—取舍—验证”的顺序表达：",
    )
    add_numbered_steps(doc, [
        "问题：普通系统提醒容易被忽略，希望利用熟人真实喝水行为形成更自然的轻提示。",
        "核心决策：把“我喝了”与“跟一口”拆成独立事件，既保留社交反馈，又阻断通知连锁。",
        "关系取舍：新增成员采用全员同意会增加步骤，但符合小规模、平等关系的产品定位。",
        "社交取舍：不显示已读和未回应，放弃监督能力，换取更平静、更低压力的使用体验。",
        "选择权设计：默认一次回应全部，通过次级入口提供选择性回应，平衡效率与微妙的人际边界。",
        "验证计划：先用 3～5 名真实朋友进行约一周测试，观察是否愿意自然发起、是否被打扰、是否理解两种行为，再迭代 V0.2。",
    ], decimal_abstract)
    add_callout(
        doc,
        "面试可提炼的一句话",
        "我没有把喝水提醒做成任务管理，而是围绕低压力熟人关系，设计了一套没有监督、没有回复债务、但保留个人边界的轻社交闭环。",
        LIGHT_BLUE_FILL,
    )

    add_heading(doc, "10. 版本记录", 1)
    add_rich_para(doc, [("V0.1 · 2026-08-17：", True, DEEP_TEAL), ("汇总核心闭环、关系建立、多信号处理、可见性控制、关系退出及离线通知流程；全部基于截至本阶段已确认的产品决策。", False, INK)])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
