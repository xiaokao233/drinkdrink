from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "跟一口_产品决策记录_QA_V0.1.docx"

# Resolved preset: decision_memo (standard_business_brief alias)
# Named override: w:eastAsia uses Microsoft YaHei for reliable Chinese rendering.
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1)
HEADER_DISTANCE = Inches(0.492)
FOOTER_DISTANCE = Inches(0.492)
FONT_LATIN = "Arial"
FONT_CJK = "Microsoft YaHei"
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)
MUTED = RGBColor(0x78, 0x78, 0x78)
LIGHT_FILL = "F4F6F9"
LIGHT_BLUE_FILL = "E8EEF5"
RULE = "B7C5D4"


DECISIONS = [
    {
        "id": "D001",
        "question": "当前阶段是否还需要重点验证“好友行为会影响喝水”这一理论？",
        "answer": "不需要把理论成立与否作为当前阶段的主要任务。已有论文理论和实际产品案例可以支持这一机制具有合理性；后续真实使用更关注具体产品体验是否自然、是否有人愿意持续使用。",
        "reason": "项目目标是完成一个可用的软件小产品，而不是进行精确的学术因果研究。",
        "impact": "核心行为模型阶段直接定义产品规则；一周测试仍保留，但定位为产品体验与真实使用观察。",
    },
    {
        "id": "D002",
        "question": "产品支持哪类关系结构？",
        "answer": "同时支持 1 对 1 和多人小组。1 对 1 可以承载情侣、搭子等关系；多人小组可以由朋友、室友或其他不同物理距离的熟人组成。",
        "reason": "两类关系共享同一个核心需求：让亲近的人真实发生的喝水行为成为自己的轻提示。",
        "impact": "后续 User Flow 应使用统一的“关系空间”思路，再分别覆盖双人和多人状态，而不是设计两套完全独立的产品。",
    },
    {
        "id": "D003",
        "question": "多人小组中谁可以发出喝水信号？",
        "answer": "小组中的任意成员都可以主动点击“我喝了”。A 发出信号时可以带动 B、C；B 发出信号时可以带动 A、C；C 发出信号时可以带动 A、B。",
        "reason": "产品模拟的是室友之间相互带动的生活行为，不设置固定的提醒者或被提醒者。",
        "impact": "成员角色是平等且可互换的，不需要创建“监督者”“被监督者”或固定领喝者身份。",
    },
    {
        "id": "D004",
        "question": "“我喝了”和“跟一口”是什么关系？",
        "answer": "两者是独立的操作语义。“我喝了”创建一条新的喝水信号；“跟一口”只回应当前信号，不创建新的喝水信号，也不再次向其他成员发出提醒。",
        "reason": "区分主动发起与轻回应，可以保留社交反馈，同时避免跟随行为形成连续通知链。",
        "impact": "数据、交互和通知都必须区分 Drink Signal 与 Follow Response；两者不能在统计和界面状态中混为一谈。",
        "key": True,
    },
    {
        "id": "D005",
        "question": "是否需要严格规定先喝水还是先点击？",
        "answer": "不需要。用户可以先喝水再点击，也可以先点击后马上喝水。产品只提供轻量行为触发，不校验现实动作的精确顺序。",
        "reason": "过度校验会增加操作负担，也超出了轻量软件产品的目标。",
        "impact": "不增加拍照、传感器验证、倒计时确认或二次提交等流程。",
    },
    {
        "id": "D006",
        "question": "一条“我喝了”信号持续多久？",
        "answer": "当前不设置明确的失效时间。它可以保持为关系空间中的当前信号，直到有成员发出下一条新的“我喝了”信号。",
        "reason": "产品不需要用倒计时制造紧迫感，最新信号机制已经足够表达“最近是谁喝了”。",
        "impact": "User Flow 不需要设计过期提示或超时失败；历史记录如何保留将在记录页阶段继续定义。",
    },
    {
        "id": "D007",
        "question": "同一成员短时间连续点击“我喝了”如何处理？",
        "answer": "短时间内连续发生的重复点击默认合并，主要用于消除误触；正常间隔后再次喝水仍然允许发出新信号。",
        "reason": "既要避免误触造成重复提示，也不能限制本身喝水频率较高的成员正常记录。",
        "impact": "防误触合并窗口的具体时长留到技术结构阶段确定，当前只确认业务原则。",
    },
    {
        "id": "D008",
        "question": "是否向没有跟随的成员显示未完成状态或再次提醒？",
        "answer": "不显示，也不再次提醒。成员没有点击“跟一口”时，系统不留下落后、缺席或未完成标签。",
        "reason": "未跟随信息容易把轻提示变成监督和任务压力，违背“不催、不卷、不监督”的产品原则。",
        "impact": "页面和通知只呈现已经发生的正向行为，不呈现谁没有跟。",
    },
    {
        "id": "D009",
        "question": "“我喝了”默认会被谁看到？",
        "answer": "默认由当前 1 对 1 关系或多人小组中的其他成员看到，不需要发起者每次选择接收对象。",
        "reason": "成员通过邀请码主动建立关系，已经具备足够的熟悉度；每次选择对象会破坏一键记录的低成本体验。",
        "impact": "“我喝了”保持单步操作，MVP 不增加逐次可见范围选择器。",
    },
    {
        "id": "D010",
        "question": "MVP 是否需要复杂的喝水隐私设置？",
        "answer": "不需要。喝水行为本身不是本项目中的高敏感信息，且关系由邀请码建立。",
        "reason": "过细的隐私层级会增加设置和理解成本，却不会提升当前核心体验。",
        "impact": "MVP 暂不设计单条信号隐藏、分组可见或临时隐身等功能。账户安全、关系退出等基础能力仍在后续技术阶段单独考虑。",
    },
    {
        "id": "D011",
        "question": "成员是否必须参与“跟一口”？",
        "answer": "不必须。成员可以不跟，也可以选择屏蔽某位成员的信号；不参与不会产生任何惩罚或负面展示。",
        "reason": "不同成员有自己的饮水习惯和节奏，产品只提供可被接受的行为提示。",
        "impact": "需要保留通知控制能力，但其性质是打扰控制而不是饮水隐私；具体入口和粒度留到页面与状态设计阶段。",
    },
    {
        "id": "D012",
        "question": "当前是否要确定喝水提示的具体视觉和反馈方式？",
        "answer": "暂不确定。当前只确认：一位成员点击“我喝了”后，其他成员应得到一次大家能够接受的轻提示。具体载体、视觉语言、动效、声音和文案由后续品牌、界面与交互反馈设计阶段决定。",
        "reason": "视觉呈现需要由产品创作者本人把控，不应在核心行为模型阶段被提前锁死。",
        "impact": "User Flow 只标注“成员感知到喝水信号”，不预设水波、弹窗、Widget 或其他具体界面表现。",
    },
    {
        "id": "D013",
        "question": "发起者是否获得他人跟随的反馈？",
        "answer": "获得轻量正向反馈。成员点击“跟一口”后，当前信号的发起者可以知道有人跟上了，但该回应不会再成为新的喝水提醒。",
        "reason": "回应能够让主动记录者感到自己的自然行为确实带动了朋友，同时不制造新的通知接力。",
        "impact": "后续状态清单需要覆盖“我发起后有人跟上”和“我跟上别人”两个不同视角。",
    },
    {
        "id": "D014",
        "question": "同一用户处于多份关系时，“我喝了”默认发给谁？",
        "answer": "默认分发给所有仍然具有有效连接的关联成员。同一接收者即使通过多份关系与发起者相连，同一次喝水信号也只呈现和提醒一次。",
        "reason": "喝水信号的隐私敏感度较低，且关系均由成员主动建立；每次发送前重新选择收件人会破坏一键表达的低成本体验。",
        "impact": "系统需要按用户去重分发，同时保留关系来源，便于后续区分不同关系中的跟随反馈。",
    },
    {
        "id": "D015",
        "question": "多名成员先后发出信号时，接收者如何处理？",
        "answer": "新的喝水信号汇入同一个当前轮次，最新信号更突出，其他来源保留为轻量背景，不形成多条未读任务。",
        "reason": "真实社交中有人会全部回应、有人只回应最近一条；产品应保留信息来源，但不能制造回复债务。",
        "impact": "首页需要覆盖单人来信与多人汇总状态；完成一次明确选择后，整个当前轮次结束。",
    },
    {
        "id": "D016",
        "question": "“跟一口”默认回应谁，能否只回应部分成员？",
        "answer": "默认一次回应当前轮次中的全部发起者；同时提供一个次级入口，让用户在少数特殊情况下只回应选中的成员。未被选中的成员不会知道自己被排除。",
        "reason": "默认路径应足够轻，但冷战或关系微妙等特殊情况也需要低负担的个人选择权。",
        "impact": "页面／状态清单需要包含默认跟随、进入选择状态、选择部分成员和选择完成四类状态；具体采用长按还是显式入口留到低保真阶段。",
    },
    {
        "id": "D017",
        "question": "用户忽略当前来信后改为点击“我喝了”，原轮次如何处理？",
        "answer": "用户可以根据自己的个人判断独立选择“我喝了”。新的主动信号正常发出，原来信轮次随之结束，之后不再补回该轮的“跟一口”。",
        "reason": "产品记录的是用户主动表达的社交意图，不需要理性界定这口水究竟受谁影响。",
        "impact": "“我喝了”既是主动发起，也是结束当前来信轮次的一种明确选择；历史来信不进入待回复列表。",
    },
    {
        "id": "D018",
        "question": "关系成员如何控制彼此的信号往来？",
        "answer": "保留“不看 TA 的信号”和“不让 TA 看我的信号”两种相互独立、可逆且静默生效的控制；也可以暂停整份关系的信号往来。",
        "reason": "人际边界可能是非对称的，类似“不看对方朋友圈”和“不让对方看我的朋友圈”，需要给用户留下微妙但不公开伤害关系的选择权。",
        "impact": "关系详情和成员设置需要分别表达接收控制与发送控制；设置变化不通知对方。",
    },
    {
        "id": "D019",
        "question": "小规模多人关系中是否存在群主、管理员或踢人权？",
        "answer": "不存在。新增成员必须经全部现有成员同意，候选人也必须最终接受；任何个人或成员组合都不能把其他成员踢出去，每个人只能自行退出。",
        "reason": "关系成员地位平等。由部分成员驱逐另一名成员会造成权力失衡，也可能让产品成为排斥或微缩霸凌的工具。",
        "impact": "需要设计全员确认、候选人接受、自主退出和关系自然结束状态，不设计群主转让、管理员或移除成员页面。",
    },
    {
        "id": "D020",
        "question": "关系名称与现实关系类型如何处理？",
        "answer": "关系名称只是便于识别的文字标签，不需要系统判断它是情侣、宿舍、朋友或其他现实关系；重要的是成员之间实际建立了哪些连接。",
        "reason": "同一类产品行为可以存在于多种现实关系中，预设关系类型不会改善核心体验，反而容易限制用户。",
        "impact": "创建关系时不要求选择固定关系类别；视觉区分可以使用名称、成员和关系来源，但具体方式留到后续设计阶段。",
    },
    {
        "id": "D021",
        "question": "V0.1 通过什么形式让朋友真实使用？",
        "answer": "采用手机优先、可添加到主屏幕的 PWA。通过网址、邀请链接或二维码让 3～5 位朋友安装和使用；App Store 上架不属于本轮 MVP 范围。",
        "reason": "本轮目标是尽快验证真实的多人喝水闭环，而不是先完成应用商店分发。PWA 能保留桌面图标、独立打开和通知能力，并允许快速迭代。",
        "impact": "页面／状态清单必须加入首次网页进入、添加到主屏幕、通知授权及网络异常状态；技术结构阶段需要支持实时同步和 Web Push。",
        "key": True,
    },
    {
        "id": "D022",
        "question": "何时开始在小红书或其他公开平台宣传？",
        "answer": "先完成 3～5 位真实朋友约一周的使用测试，再根据数据和反馈迭代；公开宣传放在产品已有稳定体验入口和可展示结果之后。",
        "reason": "亲友测试适合发现安装、通知和核心闭环问题；公开传播需要更稳定的体验，也需要能够讲清楚真实迭代过程。",
        "impact": "V0.1 的发布对象限定为受邀朋友；宣传渠道、内容和作品集材料在真实测试及 V0.2 复盘后制定。",
    },
]


PENDING = [
    ("P001", "小规模多人关系在 V0.1 中是否需要设置明确人数上限。", "页面 / 状态清单阶段"),
    ("P004", "重复点击的防误触合并窗口具体为多长。", "技术结构阶段"),
    ("P006", "最终产品名、品牌概念、提示的视觉语言和交互反馈。", "品牌与视觉概念阶段"),
    ("P007", "PWA 最终使用哪一种托管、域名和推送服务，以及国内网络下的稳定性方案。", "技术结构阶段"),
    ("P008", "首次使用过程中，在什么时机引导添加到主屏幕和申请通知权限。", "页面 / 状态清单阶段"),
    ("P009", "小红书或其他公开平台的宣传内容、节奏与体验入口。", "真实测试与 V0.2 复盘后"),
]


def set_run_font(run, size=None, color=BLACK, bold=None, italic=None):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, edge="bottom", color=RULE, size="10", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fallback)
    run._r.append(fld_char_2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.right_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 12, 6),
        "Heading 2": (13, BLUE, 10, 5),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = hp.add_run("跟一口 · 产品决策记录")
    set_run_font(left, size=9, color=MUTED, bold=True)
    right = hp.add_run("\tV0.1")
    set_run_font(right, size=9, color=MUTED)
    set_paragraph_border(hp, edge="bottom", color="D7DBE2", size="4", space="4")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    add_page_field(fp)


def add_metadata_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    label_run = p.add_run(f"{label}：")
    set_run_font(label_run, size=10.5, bold=True)
    value_run = p.add_run(value)
    set_run_font(value_run, size=10.5)
    return p


def add_labeled_paragraph(doc, label, text, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_with_next = keep_with_next
    label_run = p.add_run(f"{label}：")
    set_run_font(label_run, size=10.5, color=DARK_BLUE, bold=True)
    value_run = p.add_run(text)
    set_run_font(value_run, size=10.5, color=BLACK)
    return p


def add_qa(doc, item):
    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    id_run = heading.add_run(f"{item['id']} · ")
    set_run_font(id_run, size=10, color=DARK_BLUE, bold=True)
    question_run = heading.add_run(item["question"])
    set_run_font(question_run, size=13, color=BLUE, bold=True)

    answer = add_labeled_paragraph(doc, "Answer", item["answer"], keep_with_next=True)
    reason = add_labeled_paragraph(doc, "Reason", item["reason"], keep_with_next=True)
    impact = add_labeled_paragraph(doc, "Impact", item["impact"])
    impact.paragraph_format.space_after = Pt(8)

    if item.get("key"):
        for p in (answer, reason, impact):
            set_cell_shading(p, LIGHT_BLUE_FILL)


def add_pending(doc, item_id, question, checkpoint):
    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    id_run = heading.add_run(f"{item_id} · ")
    set_run_font(id_run, size=10, color=DARK_BLUE, bold=True)
    q_run = heading.add_run(question)
    set_run_font(q_run, size=13, color=BLUE, bold=True)
    status = add_labeled_paragraph(doc, "Status", "待后续确认", keep_with_next=True)
    checkpoint_p = add_labeled_paragraph(doc, "Checkpoint", checkpoint)
    checkpoint_p.paragraph_format.space_after = Pt(8)
    set_cell_shading(status, LIGHT_FILL)
    set_cell_shading(checkpoint_p, LIGHT_FILL)


def build():
    doc = Document()
    configure_document(doc)

    # memo_masthead title block
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(12)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    kr = kicker.add_run("PRODUCT DECISION LOG")
    set_run_font(kr, size=9.5, color=DARK_BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    tr = title.add_run("《跟一口》产品决策记录")
    set_run_font(tr, size=23, color=BLACK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("Question & Answer · V0.1")
    set_run_font(sr, size=14, color=GRAY)

    add_metadata_line(doc, "项目阶段", "页面／状态清单")
    add_metadata_line(doc, "文档状态", "活文档 · 随产品决策持续更新")
    add_metadata_line(doc, "建立日期", "2026 年 8 月 16 日")
    add_metadata_line(doc, "决策责任人", "产品创作者")

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(14)
    set_paragraph_border(rule, edge="bottom", color=RULE, size="10", space="6")

    summary = doc.add_paragraph()
    summary.paragraph_format.left_indent = Inches(0.18)
    summary.paragraph_format.right_indent = Inches(0.18)
    summary.paragraph_format.space_before = Pt(0)
    summary.paragraph_format.space_after = Pt(12)
    summary.paragraph_format.line_spacing = 1.15
    set_cell_shading(summary, LIGHT_FILL)
    lead = summary.add_run("当前核心结论  ")
    set_run_font(lead, size=11, color=DARK_BLUE, bold=True)
    body = summary.add_run("“我喝了”创建主动信号，“跟一口”只返回回应；V0.1 采用 PWA，先由 3～5 位朋友真实使用，再考虑公开传播。")
    set_run_font(body, size=11, color=BLACK, bold=True)

    h1 = doc.add_paragraph("1. 文档用途", style="Heading 1")
    h1.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(
        "本文件记录产品推进过程中已经确认的关键问题、最终回答、决策原因和影响范围。"
        "它与 PRD 配合使用：PRD 描述产品要解决什么问题和计划做什么；本文件解释关键规则为什么这样确定。"
    )
    p.paragraph_format.space_after = Pt(6)
    p2 = doc.add_paragraph(
        "更新原则：只把产品创作者已经确认的内容写入“已确认决策”；尚未决定的内容进入“待后续确认”，"
        "并在相应阶段再做选择。"
    )
    p2.paragraph_format.space_after = Pt(10)

    doc.add_paragraph("2. 已确认的产品决策", style="Heading 1")
    for item in DECISIONS:
        add_qa(doc, item)

    doc.add_paragraph("3. 当前核心行为模型", style="Heading 1")
    flow = doc.add_paragraph()
    flow.paragraph_format.space_before = Pt(2)
    flow.paragraph_format.space_after = Pt(9)
    flow.paragraph_format.line_spacing = 1.2
    set_cell_shading(flow, LIGHT_FILL)
    fr = flow.add_run("成员现实中喝水  →  点击“我喝了”  →  其他成员感知一次轻提示  →  自主选择跟或不跟  →  点击“跟一口”的成员只返回轻回应  →  下一条“我喝了”成为新的当前信号")
    set_run_font(fr, size=11, color=BLACK, bold=True)

    p = doc.add_paragraph(
        "该模型不预设具体页面、通知载体或视觉语言。User Flow 阶段只描述参与者、触发、系统行为和反馈；"
        "视觉呈现将在品牌、界面和交互反馈阶段由产品创作者进一步决定。"
    )
    p.paragraph_format.space_after = Pt(10)

    doc.add_paragraph("4. 待后续确认的问题", style="Heading 1")
    intro = doc.add_paragraph("以下问题不会阻碍当前核心行为模型定稿，但必须在标注的阶段前完成决策。")
    intro.paragraph_format.space_after = Pt(8)
    for pending in PENDING:
        add_pending(doc, *pending)

    doc.add_paragraph("5. 后续 Q&A 记录模板", style="Heading 1")
    template_intro = doc.add_paragraph(
        "未来每次出现会影响范围、交互或技术实现的新问题时，复制以下结构记录。面试或作品集复盘时，"
        "重点讲清楚 Question、Answer 与 Reason，而不只是展示最终界面。"
    )
    template_intro.paragraph_format.space_after = Pt(9)

    for label, value in [
        ("Question", "需要决定的具体问题是什么？"),
        ("Answer", "最终选择了什么？"),
        ("Reason", "为什么这样选择，放弃了哪些替代方案？"),
        ("Impact", "它会影响哪些用户流程、页面状态、数据或技术实现？"),
        ("Revisit Trigger", "出现什么新证据或条件时，需要重新讨论？"),
    ]:
        p = add_labeled_paragraph(doc, label, value)
        set_cell_shading(p, LIGHT_FILL)

    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(14)
    closing.paragraph_format.space_after = Pt(0)
    set_paragraph_border(closing, edge="top", color="D7DBE2", size="4", space="6")
    cr = closing.add_run("当前阶段：User Flow 已确认 → 页面／状态清单进行中")
    set_run_font(cr, size=10, color=MUTED, italic=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
